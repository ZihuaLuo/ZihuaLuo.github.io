"""Export aggregate, public-safe Non-Compete 10-K coding counts.

Usage: python scripts/build_noncompete_data.py SOURCE.xlsx OUTPUT.json
The export contains no filing identifiers, people, company names, text snippets,
or annotation comments.
"""

from __future__ import annotations

from collections import Counter
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


FEATURES = [
    "Restriction",
    "Type",
    "Position",
    "Duration",
    "Condition",
    "Non-restriction",
    "Geography",
    "Enforcement",
    "Penalty",
]

THEMES = {
    "Competition restriction": re.compile(r"non[\s-]*compet(?:e|ition|ion)|not[\s-]+to[\s-]+compete|against competition|covenants?-not-to-compete", re.I),
    "Non-solicitation": re.compile(r"solicit|recruit|no hire", re.I),
    "Confidentiality": re.compile(r"confidential|non[\s-]*disclos|trade secret", re.I),
}

PERIODS = [
    ("1994–2004", 1994, 2004),
    ("2005–2014", 2005, 2014),
    ("2015–2023", 2015, 2023),
]


def normalize_feature(value: object) -> str:
    label = str(value or "").strip().lower()
    if label in {"person", "people", "perosn"}:
        return "Person"
    if label in {"location", "geographic location", "area"}:
        return "Geography"
    if label in {"penalty", "penalty (enforcement)"}:
        return "Penalty"
    if label in {"enforcement", "enforcement (legal)"}:
        return "Enforcement"
    if re.fullmatch(r"condition(?:\s+[12])?", label):
        return "Condition"
    if label in {"excluding", "exception"}:
        return "Exclusion note"
    known = {name.lower(): name for name in FEATURES}
    return known.get(label, "Other")


def main(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    filings: dict[str, dict] = {}
    current: str | None = None
    annotation_rows = 0

    for row in workbook["Sheet1"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            current = str(row[0]).strip()
            if current in filings:
                raise ValueError(f"Repeated filing marker: {current}")
            match = re.match(r"^(\d{4})\d{4}_", current)
            if not match:
                raise ValueError(f"Cannot derive filing year: {current}")
            filings[current] = {"year": int(match.group(1)), "features": set(), "themes": set(), "annotationCount": 0}
        if current is None:
            raise ValueError("Annotation row before the first filing marker")
        annotation_rows += 1
        filings[current]["annotationCount"] += 1
        feature = normalize_feature(row[1])
        filings[current]["features"].add(feature)
        if feature == "Type" and row[2] is not None:
            text = str(row[2])
            for theme, pattern in THEMES.items():
                if pattern.search(text):
                    filings[current]["themes"].add(theme)

    records = list(filings.values())
    years = [record["year"] for record in records]
    year_counts = Counter(years)
    feature_counts = Counter(feature for record in records for feature in record["features"])
    theme_counts = Counter(theme for record in records for theme in record["themes"])
    period_counts = Counter(
        label for record in records for label, start, end in PERIODS if start <= record["year"] <= end
    )

    result = {
        "meta": {
            "source": source.name,
            "sheet": "Sheet1",
            "filings": len(records),
            "annotationRows": annotation_rows,
            "firstYear": min(years),
            "lastYear": max(years),
            "typeAnnotatedFilings": feature_counts["Type"],
            "exclusionNoteFilings": feature_counts["Exclusion note"],
        },
        "periods": [
            {
                "name": name,
                "startYear": start,
                "endYear": end,
                "filings": period_counts[name],
                "annotationRows": sum(record["annotationCount"] for record in records if start <= record["year"] <= end),
                "typeAnnotatedFilings": sum("Type" in record["features"] for record in records if start <= record["year"] <= end),
                "exclusionNoteFilings": sum("Exclusion note" in record["features"] for record in records if start <= record["year"] <= end),
            }
            for name, start, end in PERIODS
        ],
        "annual": [{"year": year, "filings": year_counts[year]} for year in range(min(years), max(years) + 1)],
        "features": [{"name": name, "filings": feature_counts[name]} for name in FEATURES],
        "typeThemes": [{"name": name, "filings": theme_counts[name]} for name in THEMES],
        "periodFeatures": [
            {"period": name, "feature": feature, "filings": sum(
                start <= record["year"] <= end and feature in record["features"] for record in records
            )}
            for name, start, end in PERIODS for feature in FEATURES
        ],
        "periodThemes": [
            {"period": name, "theme": theme, "filings": sum(
                start <= record["year"] <= end and theme in record["themes"] for record in records
            )}
            for name, start, end in PERIODS for theme in THEMES
        ],
    }

    if sum(item["filings"] for item in result["annual"]) != len(records):
        raise ValueError("Annual totals do not reconcile")
    if sum(item["filings"] for item in result["periods"]) != len(records):
        raise ValueError("Period totals do not reconcile")
    if sum(item["annotationRows"] for item in result["periods"]) != annotation_rows:
        raise ValueError("Period annotation totals do not reconcile")
    for item in result["features"]:
        period_total = sum(row["filings"] for row in result["periodFeatures"] if row["feature"] == item["name"])
        if period_total != item["filings"]:
            raise ValueError(f"Feature totals do not reconcile: {item['name']}")
    for item in result["typeThemes"]:
        period_total = sum(row["filings"] for row in result["periodThemes"] if row["theme"] == item["name"])
        if period_total != item["filings"]:
            raise ValueError(f"Theme totals do not reconcile: {item['name']}")

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result["meta"], ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Expected source workbook and output JSON paths")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
