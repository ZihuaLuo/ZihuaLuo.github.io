"""Export aggregate, public-safe charts from the Corporate Scandal workbook.

Usage: python scripts/build_corporate_scandal_data.py SOURCE.xlsx OUTPUT.json
Only counts are exported. Company names, case IDs, descriptions, and article text
remain in the source workbook.
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


def parse_date(value: object, epoch: datetime) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted if isinstance(converted, datetime) else None
        except (TypeError, ValueError):
            return None
    if isinstance(value, str):
        for format_string in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
            try:
                return datetime.strptime(value.strip(), format_string)
            except ValueError:
                pass
    return None


def main(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Reset Cases"]
    records = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        case_id = row[1]
        if case_id is None:
            continue
        category = str(row[5]).strip().replace("Sexual Harrasment", "Sexual Harassment")
        action = str(row[7]).strip() if row[7] is not None else "Unclassified"
        records.append({
            "case": case_id,
            "company": row[2],
            "category": category,
            "action": action,
            "date": parse_date(row[6], workbook.epoch),
        })

    case_categories: dict[object, str] = {}
    for record in records:
        prior = case_categories.setdefault(record["case"], record["category"])
        if prior != record["category"]:
            raise ValueError(f"Case {record['case']} has inconsistent categories")

    category_cases = Counter(case_categories.values())
    category_records = Counter(record["category"] for record in records)
    action_records = Counter(record["action"] for record in records)
    annual = Counter((record["date"].year, record["category"]) for record in records if record["date"])
    action_category = Counter((record["category"], record["action"]) for record in records)
    categories = sorted(category_cases, key=lambda name: (-category_cases[name], name))
    actions = sorted(action_records, key=lambda name: (-action_records[name], name))
    dated_years = [record["date"].year for record in records if record["date"]]
    article_rows = sum(row[0] is not None for row in workbook["Article List"].iter_rows(min_row=2, values_only=True))

    result = {
        "meta": {
            "source": source.name,
            "sheet": "Reset Cases",
            "codedRecords": len(records),
            "cases": len(case_categories),
            "companies": len({record["company"] for record in records}),
            "classifiedRecords": len(records) - action_records["Unclassified"],
            "datedRecords": len(dated_years),
            "undatedRecords": len(records) - len(dated_years),
            "firstYear": min(dated_years),
            "lastYear": max(dated_years),
            "separateArticleRows": article_rows,
        },
        "categories": [
            {"name": name, "cases": category_cases[name], "records": category_records[name]}
            for name in categories
        ],
        "actions": [{"name": name, "records": action_records[name]} for name in actions],
        "annual": [
            {"year": year, "category": category, "records": count}
            for (year, category), count in sorted(annual.items())
        ],
        "actionCategory": [
            {"category": category, "action": action, "records": count}
            for (category, action), count in sorted(action_category.items())
        ],
    }

    if sum(item["cases"] for item in result["categories"]) != result["meta"]["cases"]:
        raise ValueError("Case totals do not reconcile")
    if sum(item["records"] for item in result["categories"]) != result["meta"]["codedRecords"]:
        raise ValueError("Record totals do not reconcile")
    if sum(item["records"] for item in result["actions"]) != result["meta"]["codedRecords"]:
        raise ValueError("Action totals do not reconcile")
    if sum(item["records"] for item in result["annual"]) != result["meta"]["datedRecords"]:
        raise ValueError("Annual totals do not reconcile")

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result["meta"], ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Expected source workbook and output JSON paths")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
